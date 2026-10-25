#!/usr/bin/env python3
"""Exploratory EPFL Level-3 measured-power shape audit (protocol V2).

Fixed layout: Session cols 0 Session, 1 CCS, 2 Arrival, 3 Departure, 5 Energy;
Measurement cols 0 timestamp, 1 CCS, 2 Session, 3 P, 9 cumulative Energy.
Join only on (CCS, first minute, last minute) vs (CCS, Arrival, Departure).
Aggregate output only; no raw rows, session IDs, training, or inference.
"""
from __future__ import annotations

import argparse, hashlib, json, math
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
PROTOCOL = "docs/thesis/revision-20261022-evidence/epfl-proxy-audit-v2-protocol.md"
EXPECTED = {"Session_data.xlsx": "b886943bf86520f24f5a53870886c4c43709d265cbebdd65e1345321e22cc192",
            "Measurement_data.xlsx": "a0a489e2897ec79235ed914376ab7751ab21c68418543234a60042eaee73a89c"}
S_COLS = {"session": 0, "ccs": 1, "arrival": 2, "departure": 3, "energy": 5}
M_COLS = {"timestamp": 0, "ccs": 1, "session": 2, "power": 3, "energy": 9}


def sha256(path):
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def ccs(v):
    return "".join(str(v or "").upper().split())


def num(v):
    if v is None or isinstance(v, bool):
        return None
    try:
        x = float(str(v).replace(",", "").strip()) if isinstance(v, str) else float(v)
    except (TypeError, ValueError):
        return None
    return x if math.isfinite(x) else None


def dt(v):
    if isinstance(v, datetime):
        return v
    if isinstance(v, date):
        return datetime(v.year, v.month, v.day)
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return datetime(1899, 12, 30) + timedelta(days=float(v))
    if isinstance(v, str):
        try:
            return datetime.fromisoformat(v.strip())
        except ValueError:
            return None
    return None


def aligned(t):
    return t.second == 0 and t.microsecond == 0


def get(row, i):
    return row[i] if i < len(row) else None


def read_rows(path, names=None):
    from openpyxl import load_workbook
    wb = load_workbook(filename=path, read_only=True, data_only=True)
    if not names or not names.issubset(set(wb.sheetnames)):
        raise ValueError(f"Required sheets missing: {names}")
    sheets = [ws for ws in wb.worksheets if ws.title in names]
    info, rows = [], []
    for ws in sheets:
        header = next(ws.iter_rows(max_row=1, values_only=True))
        required = ({0: "Session", 1: "CCS", 2: "Arrival", 3: "Departure", 5: "Energy (Wh)"}
                    if ws.title == "Session_data" else
                    {0: "Date and time", 1: "CCS", 2: "Session", 3: "P (W)", 9: "Energy (Wh)"})
        if any(get(header, i) != value for i, value in required.items()):
            raise ValueError(f"Unexpected field layout: {ws.title}")
        if hasattr(ws, "reset_dimensions"): ws.reset_dimensions()
        data = [r for r in ws.iter_rows(min_row=2, values_only=True)
                if any(v not in (None, "") for v in r)]
        info.append({"sheet": ws.title, "rows": len(data)})
        rows.extend(data)
    wb.close()
    return info, rows


def load_sessions(rows, ex):
    out = []
    for r in rows:
        c, a, d = ccs(get(r, 1)), dt(get(r, 2)), dt(get(r, 3))
        if not c or a is None or d is None: ex["session_missing_key_field"] += 1; continue
        if c not in ("CCS1", "CCS2"): ex["session_unsupported_ccs"] += 1; continue
        if not aligned(a) or not aligned(d): ex["session_nonminute_timestamp"] += 1; continue
        if a >= d: ex["session_nonpositive_duration"] += 1; continue
        out.append((c, a, d, num(get(r, 5))))
    return out


def load_groups(rows, ex):
    grouped = {}
    for r in rows:
        c, sid, t = ccs(get(r, 1)), str(get(r, 2) or "").strip(), dt(get(r, 0))
        if not sid: ex["measurement_rows_without_session"] += 1; continue
        if not c or t is None: ex["measurement_missing_key_field"] += 1; continue
        if not aligned(t): ex["measurement_nonminute_timestamp"] += 1; continue
        x = grouped.setdefault((c, sid), ([], [], []))
        x[0].append(t); x[1].append(num(get(r, 3))); x[2].append(num(get(r, 9)))
    out = []
    for (c, _), x in grouped.items():
        a, d = min(x[0]), max(x[0])
        if a >= d: ex["measurement_group_nonpositive_span"] += 1; continue
        out.append((c, a, d, x[0], x[1], x[2]))
    return out


def match(sessions, groups, ex):
    sk, gk = defaultdict(list), defaultdict(list)
    for s in sessions: sk[s[:3]].append(s)
    for g in groups: gk[g[:3]].append(g)
    pairs = []
    for key in sorted(set(sk) | set(gk)):
        left, right = sk.get(key, []), gk.get(key, [])
        if len(left) > 1: ex["session_duplicate_key"] += len(left)
        if len(right) > 1: ex["measurement_group_duplicate_key"] += len(right)
        if len(left) == len(right) == 1: pairs.append((left[0], right[0]))
        elif left and not right: ex["session_unmatched_key"] += len(left)
        elif right and not left: ex["measurement_group_unmatched_key"] += len(right)
        elif left and right: ex["ambiguous_key_both_sides"] += len(left) + len(right)
    return pairs


def invalid(s, g):
    times, power = g[3], g[4]
    if any(p is None or p < 0 for p in power): return "nonfinite_or_negative_power"
    if len(set(times)) != len(times): return "duplicate_measurement_timestamp"
    if times != sorted(times): return "nonmonotonic_measurement_timestamp"
    times = sorted(times)
    if times[0] != s[1]: return "first_sample_not_arrival"
    if times[-1] != s[2]: return "last_sample_not_departure"
    if any(b - a != timedelta(minutes=1) for a, b in zip(times, times[1:])): return "non_one_minute_step"
    if len(times) != (times[-1] - times[0]).total_seconds() // 60 + 1: return "missing_minute_in_coverage"
    if (s[2] - s[1]).total_seconds() < 900: return "duration_below_15min"
    if s[3] is None or s[3] <= 0: return "session_energy_not_positive"
    return None


def compute(s, g):
    times, power, cumulative = g[3], g[4], g[5]
    order = sorted(range(len(times)), key=times.__getitem__)
    t = [times[i] for i in order]; p = [power[i] for i in order]; cc = [cumulative[i] for i in order]
    seg = [0.5 * (p[i] + p[i + 1]) * (t[i + 1] - t[i]).total_seconds() / 3600.0 for i in range(len(t) - 1)]
    integral = sum(seg)
    if integral <= 0: return None
    duration = (s[2] - s[1]).total_seconds(); n = max(1, int(math.ceil(duration / 900.0))); bins = [0.0] * n
    for i, e in enumerate(seg): bins[min(int((t[i] - s[1]).total_seconds() // 900), n - 1)] += e
    uni = [s[3] * max(0.0, min((i + 1) * 900.0, duration) - i * 900.0) / duration for i in range(n)]
    tv = 0.5 * sum(abs(a / integral - b / sum(uni)) for a, b in zip(bins, uni))
    delta = cc[-1] - cc[0] if all(v is not None for v in cc) and all(b >= a for a, b in zip(cc, cc[1:])) else None
    return tv, integral, delta


def stats(values):
    v = sorted(float(x) for x in values)
    if not v: return {"n": 0, "median": None, "q1": None, "q3": None, "iqr": None, "mean": None}
    def q(p):
        pos = (len(v) - 1) * p; lo, hi = int(pos), math.ceil(pos)
        return v[lo] if lo == hi else v[lo] + (v[hi] - v[lo]) * (pos - lo)
    q1, q3 = q(0.25), q(0.75)
    return {"n": len(v), "median": q(0.5), "q1": q1, "q3": q3, "iqr": q3 - q1, "mean": sum(v) / len(v)}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--session", type=Path, required=True)
    ap.add_argument("--measurements", type=Path, required=True)
    ap.add_argument("--out", type=Path, required=True)
    args = ap.parse_args()
    report = {
        "audit_id": "EPFL-PROXY-SHAPE-V2-20261022", "evidence_level": "EXPLORATORY",
        "protocol": {"path": PROTOCOL, "preserved_v1":
                     "docs/thesis/revision-20261021-evidence/epfl-proxy-audit-protocol.md",
                     "sha256": sha256(ROOT / PROTOCOL) if (ROOT / PROTOCOL).is_file() else None},
        "sources": {}, "structure": {},
        "join": {"key": ["CCS", "Arrival", "Departure"], "energy_in_key": False,
                 "session_id_used_as_cross_table_key": False},
        "exclusions": {}, "coverage": {"rule": "start at Arrival, 60 s steps, end at Departure"},
        "energy_diagnostics": {"note": "Closure/amplitude only; not calibration or a shape gate."},
        "shape": {"metric": "total_variation_distance", "evidence_level": "EXPLORATORY"},
        "interpretation_limits": [
            "Exploratory only; not calibration, ACN target error, model validation, or operational benefit.",
            "Naive wall-clock timestamps; not asserted UTC or confirmed DST.",
            "P meter boundary and timestamp semantics are unresolved in the source.",
            "Only two CCS ports; no ACN ten-node mapping, station graph, or 96x10x7 contract."],
    }
    ok = True
    for label, path in (("session", args.session), ("measurement", args.measurements)):
        actual = sha256(path) if path.is_file() else None
        expected = EXPECTED[path.name]
        ok = ok and actual == expected
        report["sources"][label] = {"path": path.as_posix(), "bytes": path.stat().st_size if path.is_file() else None,
                                    "sha256": actual, "expected_sha256": expected, "hash_match": actual == expected}
    if not ok:
        report["status"] = "STOPPED_SOURCE_HASH_MISMATCH"
    else:
        try:
            s_info, s_rows = read_rows(args.session, {"Session_data"})
            m_info, m_rows = read_rows(args.measurements, {"CCS1", "CCS2"})
        except (ImportError, ValueError) as exc:
            report["status"] = "STOPPED_PREREQUISITE"
            report["blocker"] = str(exc)
        else:
            report["structure"] = {"session_sheets": s_info, "session_rows": len(s_rows),
                                   "measurement_sheets": m_info, "measurement_rows": len(m_rows),
                                   "session_columns": S_COLS, "measurement_columns": M_COLS}
            ex = Counter()
            sessions = load_sessions(s_rows, ex)
            groups = load_groups(m_rows, ex)
            pairs = match(sessions, groups, ex)
            report["join"].update({"session_rows_parsed": len(sessions), "measurement_groups_parsed": len(groups),
                                   "one_to_one_matched_pairs": len(pairs),
                                   "matched_pairs_by_ccs": dict(Counter(s[0] for s, _ in pairs))})
            tvs, ratios, deltas, post, by_ccs = [], [], [], Counter(), defaultdict(list)
            for session, group in pairs:
                reason = invalid(session, group)
                if reason: post[reason] += 1; continue
                result = compute(session, group)
                if result is None: post["p_integral_not_positive"] += 1; continue
                tv, integral, delta = result
                tvs.append(tv); ratios.append(integral / session[3])
                by_ccs[session[0]].append(tv)
                if delta is not None: deltas.append(delta / session[3])
            report["exclusions"] = dict(sorted(ex.items()))
            report["exclusions"].update({"matched_" + k: v for k, v in sorted(post.items())})
            report["coverage"].update({"eligible_sessions": len(tvs), "excluded_after_matching": sum(post.values())})
            report["energy_diagnostics"].update({
                "p_integral_over_session_energy": stats(ratios),
                "cumulative_energy_delta_over_session_energy": stats(deltas),
                "sessions_with_cumulative_energy_delta": len(deltas)})
            report["shape"].update({"status": "COMPUTED" if tvs else "NOT_COMPUTED_NO_ELIGIBLE_SESSION",
                                    "equal_session": stats(tvs),
                                    "by_ccs": {c: stats(v) for c, v in sorted(by_ccs.items())}})
            report["status"] = "SHAPE_COMPUTED_EXPLORATORY" if tvs else "STOPPED_NO_ELIGIBLE_SESSIONS"

    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"status": report.get("status"), "out": args.out.as_posix(),
                      "matched_pairs": report["join"].get("one_to_one_matched_pairs"),
                      "eligible_sessions": report["coverage"].get("eligible_sessions")}, ensure_ascii=False, indent=2))
    return 0 if report.get("status") == "SHAPE_COMPUTED_EXPLORATORY" else 2


if __name__ == "__main__":
    raise SystemExit(main())
