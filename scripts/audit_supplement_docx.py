"""Map actual PDF body headings and audit the one canonical DOCX supplement."""
import argparse
import hashlib
import json
from pathlib import Path
import re
import subprocess
import tempfile
import unicodedata
import zipfile
from lxml import etree
from docx import Document
from generate_thesis_main_docx import toc_entries


def norm(s):
    return re.sub(r'\s+', '', unicodedata.normalize('NFKC', s))


def body_heading_positions(pdf, entries):
    """Exclude running headers and match wrapped headings in page order."""
    bbox = subprocess.check_output(['pdftotext', '-bbox-layout', str(pdf), '-'])
    tree = etree.fromstring(bbox)
    pages = tree.xpath('//*[local-name()="page"]')
    candidates = []
    for page_index, page in enumerate(pages):
        lines = []
        for line in page.xpath('.//*[local-name()="line"]'):
            y = float(line.get('yMin'))
            if 70 <= y < 770:
                words = line.xpath('./*[local-name()="word"]/text()')
                lines.append((y, float(line.get('xMin')), norm(' '.join(words))))
        lines.sort()
        for j, (y, x, content) in enumerate(lines):
            candidates.append((page_index, y, content))
            combined = content
            for k in range(j + 1, min(j + 3, len(lines))):
                if lines[k][0] - lines[k - 1][0] > 32:
                    break
                combined += lines[k][2]
                candidates.append((page_index, y, combined))
    first = norm(entries[0][1])
    start = next((p, y) for p, y, t in candidates if t == first)
    cursor = start
    result = {}
    for _, heading, _ in entries:
        matches = [(p, y) for p, y, t in candidates
                   if (p, y) >= cursor and t == norm(heading)]
        if not matches:
            raise AssertionError(f'Missing body heading: {heading}')
        cursor = min(matches)
        result[heading] = cursor[0]
    return start[0], result


def cell_text_segments(cell):
    """Plain text surrounding native math need not be adjacent in PDF text."""
    pieces = []
    for paragraph in cell.paragraphs:
        current = ''
        for child in paragraph._p:
            if child.tag.startswith('{http://schemas.openxmlformats.org/officeDocument/2006/math}'):
                if norm(current):
                    pieces.append(current)
                current = ''
            else:
                current += ''.join(child.xpath('.//*[local-name()="t"]/text()'))
        if norm(current):
            pieces.append(current)
    return pieces


def sync_cached_toc(docx, page_map):
    """Change only native TOC cached page runs; preserve all other ZIP parts."""
    document = Document(docx)
    count = 0
    for p in document.paragraphs:
        if p.style.name.lower().startswith('toc '):
            heading, separator, old_page = p.text.rpartition('\t')
            assert separator and heading in page_map and old_page.isdigit()
            texts = p._p.xpath('.//*[local-name()="t"]')
            assert texts[-1].text == old_page
            texts[-1].text = page_map[heading]
            count += 1
    assert count == len(page_map)
    xml = etree.tostring(document._element, encoding='UTF-8', xml_declaration=True, standalone=True)
    with tempfile.NamedTemporaryFile(dir=docx.parent, suffix='.docx', delete=False) as tmp:
        temporary = Path(tmp.name)
    try:
        with zipfile.ZipFile(docx) as original, zipfile.ZipFile(temporary, 'w') as target:
            for item in original.infolist():
                target.writestr(item, xml if item.filename == 'word/document.xml' else original.read(item.filename))
        temporary.replace(docx)
    finally:
        temporary.unlink(missing_ok=True)


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--docx', type=Path, required=True)
    parser.add_argument('--pdf', type=Path, required=True)
    parser.add_argument('--out', type=Path, required=True)
    parser.add_argument('--map-only', action='store_true')
    parser.add_argument('--sync-toc', action='store_true', help='Update only cached TOC page runs in the canonical DOCX')
    args=parser.parse_args()
    source=Path('docs/thesis/thesis-manuscript.md')
    s=source.read_text()
    entries=toc_entries(source)
    layout=subprocess.check_output(['pdftotext','-layout',str(args.pdf),'-']).decode().split('\f')
    if not layout[-1].strip(): layout.pop()
    start, positions = body_heading_positions(args.pdf, entries)
    page_map={}
    heading_pages={}
    for _,heading,_ in entries:
        i=positions[heading]
        footer=[line.strip() for line in layout[i].splitlines() if line.strip()][-1]
        if not re.fullmatch(r'\d+',footer): raise AssertionError(f'Missing Arabic footer on page {i+1}: {footer}')
        page_map[heading]=footer
        heading_pages[heading]=i+1
    args.out.mkdir(parents=True,exist_ok=True)
    (args.out/'toc-pages.json').write_text(json.dumps(page_map,ensure_ascii=False,indent=2)+'\n')
    if args.sync_toc:
        sync_cached_toc(args.docx, page_map)
    if args.map_only:
        print(json.dumps(dict(pdf_pages=len(layout),toc_entries=len(entries),body_start_physical_page=start+1),ensure_ascii=False)); return
    raw=subprocess.check_output(['pdftotext','-raw',str(args.pdf),'-']).decode().split('\f')
    normalized=[norm(p) for p in raw]
    d=Document(args.docx)
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
    with zipfile.ZipFile(args.docx) as z:
        xml=etree.fromstring(z.read('word/document.xml'))
        rels={r.get('Id'):r.get('Target') for r in etree.fromstring(z.read('word/_rels/document.xml.rels'))}
        expected=[hashlib.sha256((source.parent/p).resolve().with_suffix('.png').read_bytes()).hexdigest() for p in re.findall(r'!\[[^\]]*\]\(([^)]+)\)',s)]
        actual=[hashlib.sha256(z.read('word/'+rels[x])).hexdigest() for x in xml.xpath('//*[local-name()="blip"]/@*[local-name()="embed"]')]
        image_start=actual.index(expected[0]); assert actual[image_start:]==expected
        equations=[]
        for p in xml.xpath('//w:body/w:p',namespaces=ns):
            t=''.join(p.xpath('.//w:t/text()',namespaces=ns))
            if p.xpath('.//m:oMath',namespaces=ns) and re.fullmatch(r'\s*[（(](\d+-\d+)[）)]\s*',t):
                equations.append(re.search(r'\d+-\d+',t).group())
        assert equations==re.findall(r'\\tag\{([^}]+)\}',s), equations
        accents=len(xml.xpath('//m:acc',namespaces=ns))
        table_reports=[]
        for table in d.tables:
            previous=table._tbl.getprevious()
            while previous is not None:
                text=''.join(previous.xpath('.//*[local-name()="t"]/text()'))
                if re.match(r'表\d+-\d+\s',text): break
                previous=previous.getprevious()
            if previous is None: continue
            number=re.match(r'表\d+-\d+',text).group()
            matches=[i for i in range(start,len(raw)) if norm(text) in normalized[i]]
            if not matches: raise AssertionError(f'Missing table caption: {text}')
            i=matches[0]
            missing=[segment for row in table.rows for cell in row.cells
                     for segment in cell_text_segments(cell)
                     if norm(segment) not in normalized[i]]
            table_reports.append(dict(caption=text,number=number,physical_page=i+1,missing_cells=missing))
        captions=re.findall(r'\*\*(表\d+-\d+)\s',s)
        assert [r['number'] for r in table_reports]==captions
        mismatches={}
        toc=[]
        for p in d.paragraphs:
            if p.style.name.lower().startswith('toc '):
                heading,_,page=p.text.rpartition('\t'); toc.append((heading,page))
        assert len(toc)==len(entries)
        for heading,page in toc:
            if page_map.get(heading)!=page: mismatches[heading]=dict(cached=page,actual=page_map.get(heading))
        flags=[r for r in table_reports if r['missing_cells']]
        report=dict(pdf_pages=len(layout),toc_entries=len(entries),toc_mismatches=mismatches,
                    body_start_physical_page=start+1,heading_pages=heading_pages,
                    tables=table_reports,all_tables_same_page=not flags,
                    embedded_image_hash_matches=len(expected),native_numbered_equations=len(equations),native_accents=accents,
                    body_han=len(re.findall(r'[\u4e00-\u9fff]',s[s.index('# 第1章'):s.index('# 致  谢')])),
                    source_sha256=hashlib.sha256(source.read_bytes()).hexdigest(),docx_sha256=hashlib.sha256(args.docx.read_bytes()).hexdigest(),pdf_sha256=hashlib.sha256(args.pdf.read_bytes()).hexdigest())
        (args.out/'final-integrity.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
        print(json.dumps({k:v for k,v in report.items() if k not in ('heading_pages','tables')},ensure_ascii=False,indent=2))
        if flags or mismatches: raise AssertionError(f'Layout mismatch: tables={flags},toc={mismatches}')


if __name__=='__main__': main()
