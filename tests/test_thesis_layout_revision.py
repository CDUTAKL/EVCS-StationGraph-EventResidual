"""Regression checks for defects observed in the final rendered thesis."""
from docx import Document
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
import pytest
import unicodedata
import generate_thesis_main_docx as builder
from rebuild_docx_formulas import run_converter, NS


def test_outer_cover_spacing_preserves_text_and_compacts_only_empty_lines():
    from docx.shared import Pt
    doc = Document()
    for i in range(19):
        doc.add_paragraph('date' if i == 18 else '')
    before = [p.text for p in doc.paragraphs]
    builder.compact_outer_cover_spacing(doc)
    assert [p.text for p in doc.paragraphs] == before
    assert all(doc.paragraphs[i].paragraph_format.line_spacing == Pt(14) for i in (2, 6, 7))
    assert doc.paragraphs[18].paragraph_format.line_spacing is None


@pytest.mark.parametrize("latex,base,script_tag,script_text", [
    (r"10^{-3}", "10", "sSup", "−3"),
    (r"\mathbf H^{(1)}", "H", "sSup", "(1)"),
    (r"q_{\mathrm{stable},n}", "q", "sSub", "stable,n"),
])
def test_single_script_expression_retains_its_base(latex, base, script_tag, script_text):
    runtime = builder.MATH_NODE_PATH
    math = run_converter(latex, builder.NODE, str(builder.MATH_CONVERTER), runtime)
    scripts = math.xpath(f".//m:{script_tag}", namespaces=NS)
    assert len(scripts) == 1
    assert unicodedata.normalize("NFKC", "".join(scripts[0].xpath("./m:e//m:t/text()", namespaces=NS))) == base
    tag = "sup" if script_tag == "sSup" else "sub"
    assert "".join(scripts[0].xpath(f"./m:{tag}//m:t/text()", namespaces=NS)) == script_text


@pytest.mark.parametrize("latex,base,accent", [(r"\hat y", "y", "̂"), (r"\widetilde{\mathbf H}^{(l)}", "H", "̃")])
def test_prediction_and_refinement_accents_remain_native(latex, base, accent):
    runtime = builder.MATH_NODE_PATH
    math = run_converter(latex, builder.NODE, str(builder.MATH_CONVERTER), runtime)
    accents = math.xpath(".//m:acc", namespaces=NS)
    assert len(accents) == 1
    assert unicodedata.normalize("NFKC", "".join(accents[0].xpath("./m:e//m:t/text()", namespaces=NS))) == base
    assert accents[0].xpath("./m:accPr/m:chr/@m:val", namespaces=NS) == [accent]


def test_math_runs_have_explicit_size_instead_of_inheriting_a_script_size():
    runtime = builder.MATH_NODE_PATH
    math = run_converter(r"a_i-b_j", builder.NODE, str(builder.MATH_CONVERTER), runtime)
    runs = math.xpath(".//m:r", namespaces=NS)
    assert runs
    assert all(run.xpath("./w:rPr/w:sz/@w:val", namespaces=NS) == ["21"] for run in runs)


@pytest.mark.parametrize('latex,tag', [(r"P_{\mathrm{lim}}-R", 'sub'), (r"P^{\mathrm{lim}}-R", 'sup')])
def test_limit_name_in_a_variable_script_is_not_a_limit_operator(latex, tag):
    runtime = builder.MATH_NODE_PATH
    math = run_converter(latex, builder.NODE, str(builder.MATH_CONVERTER), runtime)
    literal = math.xpath(f".//m:{tag}/m:r[m:t='lim']", namespaces=NS)
    assert len(literal) == 1
    assert literal[0].xpath("./m:rPr/m:nor/@m:val", namespaces=NS) == ["1"]


def test_optimization_variable_remains_below_minimum_operator():
    runtime = builder.MATH_NODE_PATH
    math = run_converter(r"\min_{\{p_{j,h}\}}\sum_h p_{j,h}", builder.NODE, str(builder.MATH_CONVERTER), runtime)
    low = math.xpath('.//m:limLow', namespaces=NS)
    assert len(low) == 1
    assert low[0].xpath('./m:e//m:t/text()', namespaces=NS) == ['min']
    assert 'p' in low[0].xpath('./m:lim//m:t/text()', namespaces=NS)


@pytest.mark.parametrize("latex,script_tag,has_upper", [
    (r"\sum_{i\in I}x_i", "sSub", False),
    (r"\sum_{i=1}^{n}x_i", "sSubSup", True),
])
def test_sum_only_contains_nonempty_limit_scripts(latex, script_tag, has_upper):
    runtime = builder.MATH_NODE_PATH
    math = run_converter(latex, builder.NODE, str(builder.MATH_CONVERTER), runtime)
    operators = math.xpath(f".//m:{script_tag}[m:e//m:t='∑']", namespaces=NS)
    assert len(operators) == 1
    assert bool(operators[0].xpath("./m:sup", namespaces=NS)) == has_upper
    assert not math.xpath(".//m:sup[not(*)] | .//m:sub[not(*)]", namespaces=NS)


def test_requested_spacing_does_not_snap_to_template_grid():
    doc = Document()
    p = doc.add_paragraph("正文")
    builder.set_paragraph_format(p, WD_ALIGN_PARAGRAPH.JUSTIFY, True, line=1.25)
    snap = p._p.pPr.find(qn("w:snapToGrid"))
    assert snap is not None and snap.get(qn("w:val")) == "0"


def test_both_odd_and_even_pages_receive_page_number_fields():
    doc = Document()
    doc.settings.odd_and_even_pages_header_footer = True
    builder.configure_section(doc.sections[0], odd_header="第1章 绪论", start=1)
    for footer in (doc.sections[0].footer, doc.sections[0].even_page_footer):
        assert not footer.is_linked_to_previous
        fields = footer._element.findall(".//" + qn("w:fldSimple"))
        assert any(f.get(qn("w:instr")) == "PAGE" for f in fields)


def test_page_number_field_uses_the_required_small_font():
    doc = Document()
    p = doc.add_paragraph()
    builder.add_page_field(p)
    size = p._p.find(".//" + qn("w:sz"))
    assert size is not None and size.get(qn("w:val")) == "18"


def test_narrow_result_columns_break_between_mean_and_standard_deviation():
    doc = Document()
    table = builder.add_three_line_table(doc, [
        ["配置", "全局MAE", "事件MAE", "稳定MAE", "参数", "轮次"],
        ["mrg_equal_beta", "0.03862±0.00028", "0.3812±0.0014", "0.0336±0.0003", "32708", "17/16/13"],
    ])
    assert table.cell(1, 0).text == "mrg_\nequal_beta"
    assert table.cell(1, 1).text == "0.03862\n±0.00028"
    assert table.cell(1, 2).text == "0.3812\n±0.0014"
    assert table.cell(1, 4).text == "32708"


def test_wide_result_columns_keep_compact_values():
    doc = Document()
    table = builder.add_three_line_table(doc, [
        ["配置", "MAE"], ["mrg_equal_beta", "0.03862±0.00028"],
    ])
    assert table.cell(1, 0).text == "mrg_equal_beta"
    assert table.cell(1, 1).text == "0.03862±0.00028"


def test_short_tables_stay_together_without_dragging_the_following_note():
    doc = Document()
    table = builder.add_three_line_table(doc, [
        ["配置", "MAE"], ["完整模型", "0.01376"], ["无图模型", "0.02952"],
    ])
    for row in table.rows[:-1]:
        assert all(p.paragraph_format.keep_with_next for cell in row.cells for p in cell.paragraphs)
    assert all(p.paragraph_format.keep_with_next is False for cell in table.rows[-1].cells for p in cell.paragraphs)


def _front_matter_doc(tmp_path):
    source = tmp_path / "source.md"
    source.write_text("# 摘  要\n\n标签`target_event`。\n\n# Abstract\n\nUse `rule_trigger`.\n\n# 目  录\n\n# 第1章  绪论\n\n# 参考文献\n\n# 在读期间发表的学术论文及研究成果\n", encoding="utf-8")
    doc = Document()
    for level in (1, 2, 3):
        doc.styles.add_style(f"toc {level}", WD_STYLE_TYPE.PARAGRAPH)
    builder.add_front_matter(doc, source)
    return doc


def test_abstracts_do_not_show_markdown_backticks(tmp_path):
    doc = _front_matter_doc(tmp_path)
    assert not any("`" in p.text for p in doc.paragraphs)


def test_english_abstract_has_an_english_thesis_title(tmp_path):
    doc = _front_matter_doc(tmp_path)
    assert any(p.text.startswith("Short-Term Load Forecasting") for p in doc.paragraphs)
