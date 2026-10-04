"""Regression checks for wrapped TOC headings and native-math table text."""
import zipfile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from lxml import etree

from audit_supplement_docx import body_heading_positions, cell_text_segments, sync_cached_toc


def test_wrapped_body_title_wins_over_later_running_header(monkeypatch):
    bbox = b'''<html><page>
      <line yMin="45" xMin="90"><word>Chapter four</word></line>
      <line yMin="85" xMin="90"><word>Chapter</word></line>
      <line yMin="106" xMin="200"><word>four</word></line>
      <line yMin="140" xMin="90"><word>4.1 Introduction</word></line>
      </page><page>
      <line yMin="45" xMin="90"><word>Chapter four</word></line>
      <line yMin="85" xMin="90"><word>Next page</word></line>
      </page></html>'''
    monkeypatch.setattr('audit_supplement_docx.subprocess.check_output', lambda command: bbox)
    start, pages = body_heading_positions('unused.pdf', [(1, 'Chapter four', ''), (2, '4.1 Introduction', '')])
    assert start == 0 and pages == {'Chapter four': 0, '4.1 Introduction': 0}


def test_plain_text_around_math_is_checked_as_separate_segments():
    document = Document()
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    p = cell.paragraphs[0]
    p.add_run('关系级')
    p._p.append(etree.fromstring(b'<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:r><m:t>beta</m:t></m:r></m:oMath>'))
    p.add_run('门控')
    assert cell_text_segments(cell) == ['关系级', '门控']


def test_front_matter_toc_heading_is_excluded_by_body_page_filter(monkeypatch):
    bbox = b'''<html><page><line yMin="100" xMin="90"><word>Chapter one</word></line></page>
      <page><line yMin="100" xMin="90"><word>Chapter one</word></line></page></html>'''
    monkeypatch.setattr('audit_supplement_docx.subprocess.check_output', lambda command: bbox)
    start, pages = body_heading_positions('unused.pdf', [(1, 'Chapter one', '')], {1})
    assert start == 1 and pages == {'Chapter one': 1}


def test_toc_page_sync_preserves_other_zip_parts_and_body(tmp_path):
    path = tmp_path / 'single.docx'
    document = Document()
    document.styles.add_style('toc 1', WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph('Chapter four\t10', style='toc 1')
    document.add_paragraph('Unchanged body')
    document.save(path)
    with zipfile.ZipFile(path) as package:
        original = {name: package.read(name) for name in package.namelist()}
    sync_cached_toc(path, {'Chapter four': '12'})
    with zipfile.ZipFile(path) as package:
        for name, content in original.items():
            if name != 'word/document.xml':
                assert package.read(name) == content
    updated = Document(path)
    assert updated.paragraphs[0].text == 'Chapter four\t12'
    assert updated.paragraphs[1].text == 'Unchanged body'
    assert list(tmp_path.iterdir()) == [path]
